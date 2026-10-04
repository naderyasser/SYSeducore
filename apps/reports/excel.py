"""
Formatted Excel (.xlsx) attendance report.

Replaces two CSV exports the desk could not use: the scanner's "summary" was
a title row and three counts, and the report page scraped its own HTML table
— one page of 25 rows, the student cell full of line breaks, and Excel
turning the date and time cells into numbers too wide for their columns,
which it shows as ``######``. The client asked for "تقرير منسق وشكله نضيف":
teacher, group, how many attended, and how many of them paid in full, paid
part, or did not pay.

Dates are written as text on purpose: a ``YYYY-MM-DD`` string reads the same
in every Excel locale and can never collapse into ``######``.
"""
from collections import Counter, defaultdict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.attendance.models import Attendance, Session
from apps.payments.models import Payment
from apps.teachers.models import WEEK_DAYS_AR

_WEEKDAYS = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']

_TITLE_FONT = Font(bold=True, size=14, color='1E3A8A')
_HEAD_FONT = Font(bold=True, color='FFFFFF')
_HEAD_FILL = PatternFill('solid', fgColor='2563EB')
_TOTAL_FILL = PatternFill('solid', fgColor='E0E7FF')
_ZEBRA_FILL = PatternFill('solid', fgColor='F8FAFC')
_THIN = Side(style='thin', color='CBD5E1')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
_RIGHT = Alignment(horizontal='right', vertical='center', wrap_text=True)

STATUS_AR = {'present': 'حاضر', 'late': 'متأخر', 'absent': 'غائب', 'exception': 'عذر'}
PAY_AR = {'paid': 'مدفوع كامل', 'partial': 'مدفوع جزئي', 'unpaid': 'غير مدفوع', 'exempt': 'معفى'}


def _payment_lookup(attendances):
    """``(student_id, session) -> payment status`` for the payment that covers
    that lesson: the student's Payment on the session's cycle, or — for a
    group billed by month — on that month. No row at all reads as unpaid."""
    student_ids = {a.student_id for a in attendances}
    group_ids = {a.session.group_id for a in attendances}
    by_cycle, by_month = {}, {}
    for p in Payment.objects.filter(student_id__in=student_ids, group_id__in=group_ids).only(
        'student_id', 'group_id', 'cycle_id', 'month', 'status', 'is_exempt', 'amount_due', 'amount_paid',
    ):
        status = 'exempt' if p.is_exempt else p.status
        entry = (status, p.amount_paid, p.amount_due)
        if p.cycle_id:
            by_cycle[(p.student_id, p.cycle_id)] = entry
        by_month[(p.student_id, p.group_id, p.month)] = entry

    def lookup(att):
        session = att.session
        if session.cycle_id and (att.student_id, session.cycle_id) in by_cycle:
            return by_cycle[(att.student_id, session.cycle_id)]
        key = (att.student_id, session.group_id, session.session_date.replace(day=1))
        return by_month.get(key, ('unpaid', 0, 0))
    return lookup


def _write_table(ws, start_row, headers, rows, widths, total=None):
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=title)
        cell.font, cell.fill, cell.alignment, cell.border = _HEAD_FONT, _HEAD_FILL, _CENTER, _BORDER
    ws.row_dimensions[start_row].height = 30
    r = start_row
    for i, values in enumerate(rows):
        r = start_row + 1 + i
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = _BORDER
            cell.alignment = _CENTER if isinstance(value, (int, float)) else _RIGHT
            if i % 2:
                cell.fill = _ZEBRA_FILL
    if total:
        r += 1
        for col, value in enumerate(total, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font, cell.fill, cell.border = Font(bold=True), _TOTAL_FILL, _BORDER
            cell.alignment = _CENTER if isinstance(value, (int, float)) else _RIGHT
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _title(ws, text, subtitle, ncols):
    ws.sheet_view.rightToLeft = True
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(row=1, column=1, value=text).font = _TITLE_FONT
    ws.cell(row=1, column=1).alignment = _CENTER
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.cell(row=2, column=1, value=subtitle).alignment = _CENTER
    ws.cell(row=2, column=1).font = Font(color='64748B')


def build_attendance_workbook(date_from, date_to, group_id=None):
    """The report as ``bytes`` (xlsx): a per-lesson summary sheet and a
    per-student detail sheet over ``[date_from, date_to]``."""
    sessions = (
        Session.objects.filter(
            session_date__gte=date_from, session_date__lte=date_to, is_cancelled=False,
        )
        .select_related('group', 'group__teacher')
        .order_by('session_date', 'group__group_name')
    )
    if group_id is not None:
        sessions = sessions.filter(group_id=group_id)
    sessions = list(sessions)

    attendances = list(
        Attendance.objects.filter(session__in=sessions)
        .select_related('student', 'session', 'session__group', 'session__group__teacher')
        .order_by('session__session_date', 'session__group__group_name', 'student__full_name')
    )
    pay_of = _payment_lookup(attendances)

    by_session = defaultdict(list)
    for att in attendances:
        by_session[att.session_id].append(att)

    period = (
        f'يوم {date_from:%Y-%m-%d}' if date_from == date_to
        else f'من {date_from:%Y-%m-%d} إلى {date_to:%Y-%m-%d}'
    )
    wb = Workbook()

    # ── Sheet 1: one row per lesson ──
    ws = wb.active
    ws.title = 'ملخص الحصص'
    headers = [
        'التاريخ', 'اليوم', 'المدرس', 'المجموعة', 'عدد الحاضرين', 'منهم متأخر',
        'الغائبين', 'مدفوع كامل', 'مدفوع جزئي', 'غير مدفوع', 'معفى',
    ]
    _title(ws, 'تقرير الحضور والدفع', f'{period} — الدفع محسوب للحاضرين', len(headers))
    rows, totals = [], Counter()
    for s in sessions:
        atts = by_session.get(s.session_id, [])
        came = [a for a in atts if a.status in ('present', 'late', 'exception')]
        pays = Counter(pay_of(a)[0] for a in came)
        row = {
            'attended': len(came),
            'late': sum(1 for a in came if a.status == 'late'),
            'absent': sum(1 for a in atts if a.status == 'absent'),
            'paid': pays['paid'], 'partial': pays['partial'],
            'unpaid': pays['unpaid'], 'exempt': pays['exempt'],
        }
        totals.update(row)
        rows.append([
            s.session_date.strftime('%Y-%m-%d'),
            WEEK_DAYS_AR.get(_WEEKDAYS[s.session_date.weekday()], ''),
            s.group.teacher.full_name if s.group.teacher_id else '—',
            s.group.group_name,
            row['attended'], row['late'], row['absent'],
            row['paid'], row['partial'], row['unpaid'], row['exempt'],
        ])
    _write_table(
        ws, 4, headers, rows,
        widths=[13, 10, 24, 30, 13, 11, 10, 12, 12, 12, 8],
        total=['الإجمالي', '', '', f'{len(sessions)} حصة', totals['attended'], totals['late'],
               totals['absent'], totals['paid'], totals['partial'], totals['unpaid'], totals['exempt']],
    )

    # ── Sheet 2: one row per student per lesson ──
    ws2 = wb.create_sheet('تفاصيل الطلاب')
    headers2 = [
        'التاريخ', 'المدرس', 'المجموعة', 'الطالب', 'الكود', 'الحضور',
        'حالة الدفع', 'المدفوع', 'المطلوب',
    ]
    _title(ws2, 'تفاصيل الحضور والدفع لكل طالب', period, len(headers2))
    detail = []
    for a in attendances:
        status, paid, due = pay_of(a)
        detail.append([
            a.session.session_date.strftime('%Y-%m-%d'),
            a.session.group.teacher.full_name if a.session.group.teacher_id else '—',
            a.session.group.group_name,
            a.student.full_name,
            a.student.student_code,
            STATUS_AR.get(a.status, a.status),
            PAY_AR.get(status, status),
            float(paid or 0),
            float(due or 0),
        ])
    _write_table(ws2, 4, headers2, detail, widths=[13, 24, 30, 30, 12, 10, 14, 11, 11])

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def build_payments_workbook(payments, subtitle=''):
    """The payment report's rows (already filtered by the page) as .xlsx —
    every row, not the one page the old button scraped off the screen."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'المدفوعات'
    headers = [
        'الطالب', 'الكود', 'المجموعة', 'المدرس', 'الدورة / الشهر', 'المطلوب',
        'المدفوع', 'المتبقي', 'الحالة', 'تاريخ الدفع',
    ]
    _title(ws, 'تقرير المدفوعات', subtitle, len(headers))
    rows, due, paid = [], 0, 0
    for p in payments.select_related('student', 'group', 'group__teacher', 'cycle'):
        status = 'exempt' if p.is_exempt else p.status
        due += float(p.amount_due or 0)
        paid += float(p.amount_paid or 0)
        rows.append([
            p.student.full_name,
            p.student.student_code,
            p.group.group_name,
            p.group.teacher.full_name if p.group.teacher_id else '—',
            f'دورة {p.cycle.index}' if p.cycle_id else p.month.strftime('%Y-%m'),
            float(p.amount_due or 0),
            float(p.amount_paid or 0),
            float((p.amount_due or 0) - (p.amount_paid or 0)),
            PAY_AR.get(status, status),
            p.paid_on.strftime('%Y-%m-%d') if p.paid_on else '',
        ])
    _write_table(
        ws, 4, headers, rows, widths=[28, 12, 28, 22, 14, 11, 11, 11, 13, 13],
        total=['الإجمالي', f'{len(rows)} سجل', '', '', '', due, paid, due - paid, '', ''],
    )
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def build_cycle_register_workbook(group, cycles_data):
    """كشف الدورة بشكل الكشف الورقي — ورقة لكل دورة. ``cycles_data`` هو
    ``[(cycle, columns, rows), ...]`` كما يرجعه ``_cycle_register_data``."""
    wb = Workbook()
    wb.remove(wb.active)
    thick = Side(style='medium', color='000000')
    border = Border(left=thick, right=thick, top=thick, bottom=thick)
    red = Font(color='B91C1C', bold=True)
    for cycle, columns, rows in cycles_data:
        ws = wb.create_sheet(f'دورة {cycle.index}'[:31])
        headers = ['م', 'ت', 'الاسم'] + [str(c['seq']) for c in columns] + ['رقم الهاتف', 'ملاحظات']
        ncols = len(headers)
        span = f"من {cycle.started_on:%d/%m/%Y}" if cycle.started_on else ''
        if cycle.closed_on:
            span += f" إلى {cycle.closed_on:%d/%m/%Y}"
        _title(ws, f'{group.group_name} — دورة {cycle.index}', f'{group.teacher.full_name}  ·  {span}', ncols)
        dates_row = ['', 'تاريخ الدفع', ''] + [
            f"{c['session'].session_date.month}/{c['session'].session_date.day}" if c['session'] else ''
            for c in columns
        ] + ['', '']
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col, value=title)
            cell.font, cell.alignment, cell.border = Font(bold=True, size=12), _CENTER, border
            sub = ws.cell(row=5, column=col, value=dates_row[col - 1])
            sub.font, sub.alignment, sub.border = Font(color='64748B', size=9), _CENTER, border
        for i, r in enumerate(rows, start=1):
            values = [i, r['pay_text'], r['student'].full_name] + [c['text'] for c in r['cells']] + [r['phone'], r['note']]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=5 + i, column=col, value=value)
                cell.border = border
                cell.alignment = _RIGHT if col == 3 else _CENTER
                if value == 'غ' or (col == ncols and r['pay_status'] in ('unpaid', 'partial')):
                    cell.font = red
            ws.row_dimensions[5 + i].height = 24
        for extra in range(3):
            for col in range(1, ncols + 1):
                ws.cell(row=6 + len(rows) + extra, column=col).border = border
        widths = [5, 10, 26] + [8] * len(columns) + [15, 16]
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = 'D6'
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
